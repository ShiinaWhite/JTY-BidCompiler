"""跨项目污染检查（GPT 1.5.1 第 9 项）。

公司素材库是跨项目共享的，但**业务结论**不该跨项目串味：

* A 项目的矩阵/待办里，不许冒出 B 项目的项目变量（项目名、编号、甲方、
  投标截止日、特征设备词、保证金金额——v1.5 真实事故）；
* 反过来亦然。

扫描范围刻意**排除素材本体与证据**（source_file / source_ref / evidence_spans /
quote 等）——素材是公司级资产，证据引用本来就该指向原始出处。只扫**编译器自己
生成的文字**：reason、missing_info、note、condition_text、detail、action、summary。

金额类 token 必须写成带边界断言的正则，避免大金额误命中其中的小金额子串。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

#: 项目变量词表**只来自 config/<profile>.json**，代码里不放任何真实项目词。
#: 每个 token 可以是字符串，也可以是 {"pattern": <正则>, "label": <说明>}（金额类）。
#: 具体词表见 docs/repository_publication_audit.md "配置模板" 一节。

#: 允许出现历史引用的字段名（素材出处类）
ALLOWED_KEYS = {
    "source_file", "source_ref", "source_pages", "sources", "source",
    "provenance", "evidence_spans", "evidence_refs", "evidence",
    "quote", "source_text", "source_quote", "title", "selected_title",
    "fixture_path", "path", "targets",
}

#: 参与扫描的产物（GPT 1.5.2 第 2 项：最终生成的 JSON、MD、XLSX 都要扫）。
#: materials.json 是公司级素材库，明确排除。
SCAN_GLOBS = ("*.json", "*.md", "*.xlsx")

_TEXT_KEYS = {"reason", "missing_info", "note", "notes", "condition_text",
              "detail", "action", "summary", "label", "verdict_text"}


def _iter_text_nodes(node, path: str = "", key: str = ""):
    """遍历 JSON，产出 (路径, 文本)。允许字段与素材本体跳过。

    key 单独传递：列表内的元素路径会带 "[i]" 后缀，直接用路径尾段判断会把
    missing_info[0] 这类**列表内的文本**全部漏掉（实测漏报）。
    """
    if isinstance(node, dict):
        for k, value in node.items():
            if k in ALLOWED_KEYS:
                continue
            child_path = f"{path}.{k}" if path else k
            yield from _iter_text_nodes(value, child_path, key=k)
    elif isinstance(node, list):
        for i, item in enumerate(node):
            yield from _iter_text_nodes(item, f"{path}[{i}]", key=key)
    elif isinstance(node, str):
        if key in _TEXT_KEYS:
            yield path, node


def _token_hits(text: str, token) -> bool:
    if isinstance(token, str):
        return token in text
    return bool(re.search(token["pattern"], text))


def _iter_xlsx_text(path: Path):
    """Excel 逐单元格产出 (坐标, 文本)。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip():
                        yield f"{ws.title}!{cell.coordinate}", str(cell.value)
    finally:
        wb.close()


def scan_directory(run_dir: str | Path, *, forbidden_tokens: list,
                   forbidden_label: str) -> dict:
    """扫描一次运行目录的产物（JSON + MD + XLSX），返回命中列表。"""
    d = Path(run_dir)
    hits: list[dict] = []
    scanned = 0

    files: list[Path] = []
    for pattern in SCAN_GLOBS:
        files.extend(d.glob(pattern))
    # contamination.json 自己必然引用禁用词（自引用）；
    # materials.json 是公司级素材库的输入副本，其 notes 本就该记录跨项目出处——
    # 两者都不是"编译器生成的业务结论"，不参与扫描。
    excluded = {"contamination.json", "materials.json", "fixture_manifest.json"}
    files = sorted(f for f in set(files) if f.name not in excluded)

    for f in files:
        scanned += 1
        texts = []
        if f.suffix.lower() == ".json":
            try:
                data = json.loads(f.read_text(encoding="utf-8"))
                texts.extend(_iter_text_nodes(data, f.name))
            except Exception as exc:  # noqa: BLE001
                hits.append({"file": f.name, "path": "", "token": "<parse>",
                             "snippet": str(exc)[:120]})
        elif f.suffix.lower() in (".md", ".txt"):
            texts.append((f.name, f.read_text(encoding="utf-8", errors="replace")))
        elif f.suffix.lower() == ".xlsx":
            try:
                texts.extend(_iter_xlsx_text(f))
            except Exception as exc:  # noqa: BLE001
                hits.append({"file": f.name, "path": "", "token": "<xlsx>",
                             "snippet": str(exc)[:120]})
        for path, text in texts:
            for token in forbidden_tokens:
                if _token_hits(text, token):
                    label = token if isinstance(token, str) else token.get("label", token["pattern"])
                    hits.append({
                        "file": f.name,
                        "path": path,
                        "token": label,
                        "snippet": text[:140],
                    })
    return {
        "forbidden_project": forbidden_label,
        "files_scanned": scanned,
        "hit_count": len(hits),
        "hits": hits,
        "clean": not hits,
    }


def load_profile_tokens(profile: str) -> tuple[list, str] | None:
    """取某项目 profile 的禁用词（即其他项目的变量）。"""
    here = Path(__file__).resolve()
    for parent in [here.parent, *here.parents]:
        cfg = parent / "config" / f"{profile}.json"
        if cfg.exists():
            data = json.loads(cfg.read_text(encoding="utf-8"))
            cont = data.get("contamination") or {}
            tokens = cont.get("forbidden_tokens")
            if tokens is None:
                return None
            return tokens, cont.get("forbidden_project", "")
            break
    return None
