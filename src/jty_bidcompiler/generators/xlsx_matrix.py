"""《投标要求矩阵.xlsx》生成（MVP-2 产物，MVP-4 回填判定）。

列（按用户口径）：编号 / 招标章节 / 要求原文 / 要求类型 / 重要性 / 满足条件 /
需提供证据 / 当前候选材料 / 判定 / 缺失信息 / 来源页码 / 备注。

设计取向：公司人员打开这张表，5 秒内要能回答三件事——
"什么必须做""现在有没有""还缺什么"。
所以：HARD 置顶、判定列上色、缺失信息合并成一句话。
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..matchers.engine import VERDICT_LEXICON

HEADERS = [
    ("编号", 16), ("招标章节", 20), ("要求原文", 60), ("要求类型", 14),
    ("重要性", 9), ("满足条件", 40), ("需提供证据", 22),
    ("当前候选材料", 34), ("判定", 16), ("缺失信息", 36), ("来源页码", 10), ("备注", 26),
]

_SEVERITY_ORDER = {"HARD": 0, "SCORE": 1, "NORMAL": 2, "INFO": 3}
_SEVERITY_CN = {"HARD": "★硬性", "SCORE": "评分", "NORMAL": "一般", "INFO": "信息"}

_STATUS_FILL = {
    "PASS": PatternFill("solid", fgColor="D9EAD3"),
    "FAIL": PatternFill("solid", fgColor="F4CCCC"),
    "EVIDENCE_INSUFFICIENT": PatternFill("solid", fgColor="FFF2CC"),
    "WAIT_COMPANY": PatternFill("solid", fgColor="FCE5CD"),
    "NOT_EVALUATED": PatternFill("solid", fgColor="EFEFEF"),
}
_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
_THIN = Side(style="thin", color="B7B7B7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def write_requirements_matrix(
    requirements: list[dict],
    matrix: Optional[dict],
    out_path: str | Path,
    *,
    project_title: str = "",
) -> Path:
    entries = {e["requirement_id"]: e for e in (matrix or {}).get("entries", [])}
    wb = Workbook()
    ws = wb.active
    ws.title = "投标要求矩阵"

    _title_row(ws, f"{project_title} —— 投标要求矩阵", len(HEADERS))
    ws.append([h for h, _ in HEADERS])
    _style_header(ws, len(HEADERS))

    rows = sorted(
        requirements,
        key=lambda r: (_SEVERITY_ORDER.get(r.get("severity", "NORMAL"), 9), r.get("source_page", 0)),
    )
    for req in rows:
        entry = entries.get(req["id"], {})
        status = entry.get("status", req.get("status", "NOT_EVALUATED"))
        ws.append([
            req["id"],
            req.get("chapter", ""),
            _clip(req.get("source_text", ""), 900),
            req.get("requirement_type", ""),
            _SEVERITY_CN.get(req.get("severity", ""), req.get("severity", "")),
            _clip(req.get("condition_text") or _condition_text(req.get("condition")), 300),
            "、".join(req.get("evidence_required") or []) or "—",
            _candidates_text(entry),
            f"{status} {VERDICT_LEXICON.get(status, '')}".strip(),
            _clip("；".join(entry.get("missing_info") or []), 300) or "—",
            req.get("source_page", ""),
            _clip(_note(entry, req), 300),
        ])
    _style_body(ws, len(HEADERS), start_row=3, status_col=9)

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # 第二个 sheet：判定汇总
    summary = wb.create_sheet("判定汇总")
    summary.append(["判定", "含义", "条数"])
    for k, v in (matrix or {}).get("summary", {}).items():
        summary.append([k, VERDICT_LEXICON.get(k, ""), v])
    _style_header(summary, 3, header_row=1)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 10

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
def _title_row(ws, text: str, ncols: int) -> None:
    ws.append([text])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1)
    cell.font = Font(bold=True, size=13)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26


def _style_header(ws, ncols: int, header_row: int = 2) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[header_row].height = 22


def _style_body(ws, ncols: int, *, start_row: int, status_col: int) -> None:
    for i, (header, width) in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in range(start_row, ws.max_row + 1):
        ws.row_dimensions[r].height = None
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (3, 6, 8, 10, 12)))
            cell.border = _BORDER
            cell.font = Font(size=9)
        status_cell = ws.cell(row=r, column=status_col)
        key = str(status_cell.value or "").split(" ")[0]
        if key in _STATUS_FILL:
            status_cell.fill = _STATUS_FILL[key]
        status_cell.alignment = Alignment(vertical="top", horizontal="center", wrap_text=True)


def _candidates_text(entry: dict) -> str:
    if not entry:
        return "—"
    parts = []
    for cand in entry.get("candidates") or []:
        mid = cand.get("material_id", "")
        mark = {"PASS": "✓", "FAIL": "✗"}.get(cand.get("status", ""), "·")
        if entry.get("selected") and mid == entry["selected"]:
            mark = "★"
        reasons = "；".join(cand.get("reasons") or [])[:60]
        parts.append(f"{mark}{mid}" + (f"（{reasons}）" if reasons else ""))
    return _clip("  ".join(parts), 400) or "—"


def _note(entry: dict, req: dict) -> str:
    bits = []
    if entry.get("reason"):
        bits.append(entry["reason"])
    if entry.get("owner"):
        bits.append(f"责任方={entry['owner']}")
    if req.get("note"):
        bits.append(req["note"])
    return " ｜ ".join(bits)


def _condition_text(cond) -> str:
    if not cond:
        return ""
    if isinstance(cond, dict):
        return "；".join(f"{k}={v}" for k, v in cond.items() if k != "rule")
    return str(cond)


def _clip(text: str, limit: int) -> str:
    from ..io.textnorm import excel_safe

    text = excel_safe(text).strip()
    return text if len(text) <= limit else text[:limit] + "…"
